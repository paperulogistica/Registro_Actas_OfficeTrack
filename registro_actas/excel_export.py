"""Genera el Excel de salida con el mismo diseño de la maqueta aprobada."""
import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .parser import COLUMNS, REVIEW_COL, SOURCE_FILE_COL, PROCESSED_AT_COL

DARK_BLUE = "FF1F3864"
FLAG_FILL = "FFFFF2CC"

HEADERS = COLUMNS + [REVIEW_COL, SOURCE_FILE_COL, PROCESSED_AT_COL]

COL_WIDTHS = [10.71, 15.29, 18.43, 10.57, 15.71, 38.29, 9.57, 25.29, 10.43,
              15.43, 20.71, 15.43, 15.0, 10.86, 13.0, 14.86, 16.71, 10.57,
              12.86, 23.57, 13.57, 40.0, 30.0, 18.0]


def build_excel_bytes(all_rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Registro de Actas"

    thin = Side(style="thin", color="FFBFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    n_cols = len(HEADERS)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    title_cell = ws.cell(row=1, column=1, value="Registro de Actas de Entrega / Devolución de Equipos")
    title_cell.font = Font(bold=True, size=8, color=DARK_BLUE, name="Arial")
    title_cell.alignment = Alignment(horizontal="center")
    ws.row_dimensions[2].height = 3.75

    header_row = 3
    for idx, h in enumerate(HEADERS, start=1):
        c = ws.cell(row=header_row, column=idx, value=h)
        c.font = Font(bold=True, size=8, color="FFFFFFFF", name="Arial")
        c.fill = PatternFill("solid", fgColor=DARK_BLUE)
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        c.border = border
    ws.row_dimensions[header_row].height = 33.75
    ws.freeze_panes = "A4"

    for i, width in enumerate(COL_WIDTHS[:n_cols], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    r = header_row + 1
    for row in all_rows:
        flagged = bool(row.get(REVIEW_COL))
        for idx, col_name in enumerate(HEADERS, start=1):
            val = row.get(col_name, "")
            c = ws.cell(row=r, column=idx, value=val if val != "" else None)
            c.font = Font(size=8, name="Arial")
            c.border = border
            if flagged:
                c.fill = PatternFill("solid", fgColor=FLAG_FILL)
        r += 1

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
